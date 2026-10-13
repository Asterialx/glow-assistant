from pathlib import Path
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

root = Path(__file__).resolve().parents[1] / "tests" / "fixtures" / "office"
root.mkdir(parents=True, exist_ok=True)
doc = Document()
doc.add_heading("Исследование Glow", 0)
paragraph = doc.add_paragraph()
paragraph.add_run("Контрольная ")
paragraph.add_run("группа").bold = True
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = "Group"
table.cell(0, 1).text = "Patients"
table.cell(1, 0).text = "Placebo"
table.cell(1, 1).text = "42"
doc.save(root / "sample.docx")
wb = Workbook()
sheet = wb.active
sheet.title = "Trials"
sheet.append(["Group", "Patients"])
sheet.append(["Placebo", 42])
sheet["C2"] = "=B2*2"
second = wb.create_sheet("Notes")
second["B3"] = "Источник проверен"
wb.save(root / "sample.xlsx")
slides = Presentation()
for title, body in [("Glow Research", "Первичный исход"), ("Limitations", "Не является рекомендацией лечения")]:
    slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = body
slides.save(root / "sample.pptx")
print("Created real DOCX, XLSX and PPTX fixtures")
